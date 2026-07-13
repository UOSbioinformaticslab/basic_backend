from fastapi import APIRouter, File, UploadFile, Form, HTTPException
from typing import Optional
import json
import os
from pydantic import ValidationError
from google import genai
from google.genai import types
from ai_schemas import ExtractionResponse
from helpers.ai_prompts import SYSTEM_PROMPT
import PyPDF2
import io

router = APIRouter(
    prefix="/api",
    tags=["AI Extraction"]
)

@router.post("/extract-metadata", response_model=ExtractionResponse)
async def extract_metadata(
    file: UploadFile = File(...),
    current_form_data: Optional[str] = Form(None)
):
    filename = file.filename.lower()
    allowed_extensions = ('.pdf', '.docx', '.pptx', '.txt', '.xls', '.xlsx')
    if not filename.endswith(allowed_extensions):
        raise HTTPException(status_code=400, detail="Only PDF, DOCX, PPTX, TXT, and XLS/XLSX files are supported.")

    file_content = await file.read()
    text_content = ""

    try:
        if filename.endswith('.pdf'):
            import PyPDF2
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            num_pages = len(pdf_reader.pages)
            if num_pages > 20:
                raise HTTPException(status_code=400, detail=f"Document exceeds the 20-page limit. The uploaded document has {num_pages} pages.")
            for page in pdf_reader.pages:
                text_content += page.extract_text() + "\n"
        
        elif filename.endswith('.docx'):
            import docx
            doc = docx.Document(io.BytesIO(file_content))
            for para in doc.paragraphs:
                text_content += para.text + "\n"
                
        elif filename.endswith('.pptx'):
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_content))
            if len(prs.slides) > 20:
                raise HTTPException(status_code=400, detail=f"Document exceeds the 20-slide limit. The uploaded document has {len(prs.slides)} slides.")
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_content += shape.text + "\n"
                        
        elif filename.endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    if row_text.strip():
                        text_content += row_text + "\n"

        elif filename.endswith('.xls'):
            import xlrd
            wb = xlrd.open_workbook(file_contents=file_content)
            for sheet in wb.sheets():
                for row_idx in range(sheet.nrows):
                    row_text = " ".join([str(cell.value) for cell in sheet.row(row_idx) if cell.value is not None and str(cell.value).strip() != ""])
                    if row_text.strip():
                        text_content += row_text + "\n"

        elif filename.endswith('.txt'):
            text_content = file_content.decode('utf-8', errors='ignore')

    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(status_code=400, detail=f"Failed to read {filename.split('.')[-1].upper()} file: {str(e)}")

    # Parse current_form_data if provided
    parsed_form_data = {}
    if current_form_data:
        try:
            parsed_form_data = json.loads(current_form_data)
        except json.JSONDecodeError:
            raise HTTPException(status_code=400, detail="current_form_data must be valid JSON.")

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY environment variable is not set.")

    # Initialize Gemini client
    client = genai.Client(api_key=api_key)

    prompt = f"""
    {SYSTEM_PROMPT}

    ---
    CURRENT FORM DATA:
    {json.dumps(parsed_form_data, indent=2)}
    
    ---
    DOCUMENT TEXT:
    {text_content}
    """

    try:
        response = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=ExtractionResponse,
            ),
        )
        
        # The response.text should be a JSON string matching ExtractionResponse
        result_json = json.loads(response.text)
        return ExtractionResponse(**result_json)
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gemini API error: {str(e)}")
